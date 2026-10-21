# -*- coding: utf-8 -*-
"""
Reporte de operacion (Last Mile) -> Google Sheets

Baja de Reportes de operacion (envios.adminml.com/carriers/reports,
pestaña General, Last Mile) el mismo Excel que da el boton "Descargar CSV",
para el rango de fechas que se elija, y lo manda al Google Sheets
"Reportes de operacion ML" por su Web App:

  Extraccion   lo de esta corrida
  Consolidado  una fila por Fecha + Id de la ruta, en orden cronologico
  Registro     que se cargo y cuando

El boton de la pagina pide:
  /api/carriers/reports?mile=LM&init_date=AAAA-MM-DD&end_date=AAAA-MM-DD
                       &report_type=carrier
y la respuesta es el .xlsx (hoja "Metricas": una fila por ruta y por dia).
Aqui se pide igual, con fetch desde la pagina: la sesion la pone Chrome.

Se usa un perfil de Chrome propio (chrome_reportes): la sesion queda
guardada y la siguiente vez casi nunca pide entrar de nuevo.

Uso sin ventana:
    python extraer_reporte_operacion.py 2026-10-01 2026-10-15
"""

import base64
import io
import json
import os
import sys
import threading
import time
import urllib.error
import urllib.parse
import urllib.request
from datetime import date, datetime, timedelta

from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.chrome.service import Service

if getattr(sys, "frozen", False):
    BASE_DIR = os.path.dirname(sys.executable)
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))

PROFILE_DIR = os.path.join(BASE_DIR, "chrome_reportes")
NOMBRE_CONFIG = "webapp_reportes.json"

URL_PAGINA = "https://envios.adminml.com/carriers/reports"
URL_API = "https://envios.adminml.com/api/carriers/reports"

# El panel deja elegir un mes; se pide en tramos de hasta 31 dias
DIAS_POR_TRAMO = 31
# Filas por envio a la Web App (cada una pesa ~250 bytes)
FILAS_POR_ENVIO = 2500
# Minutos esperando a que la persona inicie sesion
ESPERA_LOGIN = 10

# La ventana se cuelga aqui para recibir los mensajes del registro
AVISAR = None


def log(msg):
    print(f"[{datetime.now():%H:%M:%S}] {msg}", flush=True)
    if AVISAR:
        try:
            AVISAR(msg)
        except Exception:
            pass


# ---------------------------------------------------------------- Google

class WebApp:
    """El Google Sheets, visto a traves de su Web App (ver
    apps_script/reportes_operacion.gs). El script corre en la cuenta del
    dueño, asi que aqui solo hacen falta la URL /exec y la clave."""

    def __init__(self, url, token):
        self.url = url
        self.token = token

    def _pedir(self, cuerpo=None):
        params = urllib.parse.urlencode({"token": self.token})
        datos, cabeceras = None, {}
        if cuerpo is not None:
            datos = json.dumps({"token": self.token, **cuerpo},
                               ensure_ascii=False).encode("utf-8")
            cabeceras["Content-Type"] = "application/json"
        # Apps Script responde con un 302 a googleusercontent.com; urllib lo
        # sigue con un GET, que es justo como hay que leer la respuesta
        req = urllib.request.Request(f"{self.url}?{params}", data=datos,
                                     headers=cabeceras,
                                     method="POST" if datos else "GET")
        try:
            with urllib.request.urlopen(req, timeout=300) as r:
                texto = r.read().decode("utf-8", "replace")
        except urllib.error.HTTPError as e:
            raise RuntimeError(f"La Web App respondio HTTP {e.code}")
        except urllib.error.URLError as e:
            raise RuntimeError(f"No se pudo conectar con la Web App: {e.reason}")
        try:
            resp = json.loads(texto)
        except json.JSONDecodeError:
            raise RuntimeError(
                "La Web App no devolvio JSON. Revisa que la URL termine en "
                "/exec y que se haya ejecutado instalar() en el script. "
                f"Empieza con: {' '.join(texto[:100].split())}")
        if not resp.get("ok"):
            if resp.get("error") == "token":
                raise RuntimeError(f"La Web App rechazo la clave de {NOMBRE_CONFIG}.")
            raise RuntimeError(f"La Web App respondio un error: {resp.get('error')}")
        return resp

    def probar(self):
        return self._pedir()

    def cargar(self, encabezados, filas, desde, hasta, reiniciar):
        return self._pedir({"accion": "cargar", "encabezados": encabezados,
                            "filas": filas, "desde": desde, "hasta": hasta,
                            "reiniciar": reiniciar})


def abrir_hoja():
    """webapp_reportes.json junto al programa manda; si no, el del .exe."""
    candidatos = [os.path.join(BASE_DIR, NOMBRE_CONFIG)]
    interno = getattr(sys, "_MEIPASS", None)
    if interno:
        candidatos.append(os.path.join(interno, NOMBRE_CONFIG))
    for ruta in candidatos:
        if os.path.isfile(ruta):
            with open(ruta, encoding="utf-8-sig") as f:
                config = json.load(f)
            url, token = config.get("url", ""), config.get("token", "")
            if "/exec" not in url or not token:
                raise RuntimeError(f"Completa url y token en {ruta}.")
            return WebApp(url, token)
    raise RuntimeError(f"Falta {NOMBRE_CONFIG} junto al programa.")


# ---------------------------------------------------------------- Chrome

def _sin_ventana():
    import subprocess
    return getattr(subprocess, "CREATE_NO_WINDOW", 0)


def limpiar_lock():
    """Si quedo un Chrome con este perfil (de una corrida anterior), se
    cierra: dos Chrome no pueden usar el mismo perfil."""
    if os.name == "nt":
        import subprocess
        marca = os.path.basename(PROFILE_DIR)
        ps = ("Get-CimInstance Win32_Process -Filter \"Name='chrome.exe'\" | "
              f"Where-Object {{ $_.CommandLine -like '*{marca}*' }} | "
              "ForEach-Object { Stop-Process -Id $_.ProcessId -Force "
              "-ErrorAction SilentlyContinue; $_.ProcessId }")
        try:
            res = subprocess.run(["powershell", "-NoProfile", "-NonInteractive",
                                  "-Command", ps], capture_output=True, text=True,
                                 timeout=25, creationflags=_sin_ventana())
            if (res.stdout or "").strip():
                log("Se cerro el Chrome de la corrida anterior.")
                time.sleep(2)
        except Exception:
            pass
    for nombre in ("lockfile", "SingletonLock", "SingletonCookie",
                   "SingletonSocket", "DevToolsActivePort"):
        try:
            os.remove(os.path.join(PROFILE_DIR, nombre))
        except OSError:
            pass


def crear_driver():
    limpiar_lock()
    opts = Options()
    opts.add_argument(f"--user-data-dir={PROFILE_DIR}")
    opts.add_argument("--profile-directory=Default")
    opts.add_argument("--start-maximized")
    opts.add_argument("--disable-blink-features=AutomationControlled")
    opts.add_argument("--lang=es-MX")
    opts.add_experimental_option("excludeSwitches", ["enable-automation"])
    # Al terminar, Chrome queda abierto para que la persona lo cierre
    opts.add_experimental_option("detach", True)
    opts.page_load_strategy = "none"
    servicio = Service(popen_kw={"creation_flags": _sin_ventana()})
    try:
        return webdriver.Chrome(options=opts, service=servicio)
    except Exception as e:
        texto = str(e).lower()
        if "cannot find chrome" in texto or "no chrome binary" in texto:
            raise RuntimeError("No se encontro Google Chrome en esta computadora.")
        raise RuntimeError(f"Chrome no pudo arrancar: {str(e).splitlines()[0]}")


def soltar_driver(driver):
    """Apaga solo el chromedriver: la ventana de Chrome sigue abierta."""
    hilo = threading.Thread(target=lambda: _stop(driver), daemon=True)
    hilo.start()
    hilo.join(20)


def _stop(driver):
    try:
        driver.service.stop()
    except Exception:
        pass


def esperar_sesion(driver, detener=None):
    """Abre la pagina de reportes y espera a que se vea con sesion."""
    driver.get(URL_PAGINA)
    log("Si Chrome pide iniciar sesion en Mercado Libre, hazlo ahi; "
        "el programa sigue solo.")
    limite = time.time() + ESPERA_LOGIN * 60
    while time.time() < limite:
        if detener and detener.is_set():
            raise RuntimeError("Detenido.")
        time.sleep(1.5)
        try:
            url = driver.current_url
        except Exception:
            continue
        if url.startswith(URL_PAGINA):
            time.sleep(2)
            return
    raise RuntimeError(f"En {ESPERA_LOGIN} minutos no se vio la pagina de "
                       "Reportes de operacion con la sesion iniciada.")


# Pide el Excel con la sesion de la pagina y lo devuelve en base64
JS_DESCARGA = r"""
const [url] = arguments;
const done = arguments[arguments.length - 1];
fetch(url, {credentials: 'include'})
  .then(r => r.arrayBuffer().then(buf => {
    const bytes = new Uint8Array(buf);
    let bin = '';
    for (let i = 0; i < bytes.length; i += 0x8000) {
      bin += String.fromCharCode.apply(null, bytes.subarray(i, i + 0x8000));
    }
    done({status: r.status, tipo: r.headers.get('content-type') || '',
          ruta: r.url, datos: btoa(bin)});
  }))
  .catch(e => done({status: 0, error: String(e)}));
"""


def descargar_tramo(driver, desde, hasta):
    """El .xlsx de ese rango, como bytes."""
    params = urllib.parse.urlencode({
        "mile": "LM", "init_date": desde.isoformat(),
        "end_date": hasta.isoformat(), "report_type": "carrier"})
    driver.set_script_timeout(300)
    for intento in range(3):
        r = driver.execute_async_script(JS_DESCARGA, f"{URL_API}?{params}") or {}
        st = r.get("status")
        if st == 200 and "json" not in (r.get("tipo") or ""):
            return base64.b64decode(r["datos"])
        if st in (0, 401, 403) or "/carriers/" not in (r.get("ruta") or URL_API):
            raise RuntimeError("La sesion de Mercado Libre caduco. Vuelve a "
                               "presionar EXTRAER e inicia sesion en Chrome.")
        detalle = ""
        if r.get("datos"):
            try:
                detalle = base64.b64decode(r["datos"]).decode("utf-8", "replace")[:150]
            except Exception:
                pass
        log(f"   El reporte respondio {st} {detalle or r.get('error', '')}; reintento...")
        time.sleep(5 * (intento + 1))
    raise RuntimeError(f"No se pudo bajar el reporte del {desde:%d/%m} al {hasta:%d/%m}.")


def leer_excel(contenido):
    """(encabezados, filas) de la hoja de metricas (la primera)."""
    from openpyxl import load_workbook

    libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    hoja = libro.worksheets[0]
    filas = list(hoja.iter_rows(values_only=True))
    if not filas:
        return [], []
    encabezados = [str(c or "").strip() for c in filas[0]]
    ancho = len(encabezados)
    while ancho and not encabezados[ancho - 1]:
        ancho -= 1
    encabezados = encabezados[:ancho]
    datos = []
    for f in filas[1:]:
        f = list(f[:ancho]) + [None] * (ancho - len(f))
        if not any(v not in (None, "") for v in f):
            continue
        datos.append([_celda(v) for v in f])
    return encabezados, datos


def _celda(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, str):
        return " ".join(v.split())
    return v


def tramos(desde, hasta):
    ini = desde
    while ini <= hasta:
        fin = min(hasta, ini + timedelta(days=DIAS_POR_TRAMO - 1))
        yield ini, fin
        ini = fin + timedelta(days=1)


def ejecutar(desde, hasta, detener=None):
    """La corrida completa. Devuelve un resumen (dict)."""
    if hasta < desde:
        desde, hasta = hasta, desde
    resumen = {"filas": 0, "nuevas": 0, "actualizadas": 0, "iguales": 0,
               "total": 0, "minutos": 0.0, "error": ""}
    inicio = time.time()
    web = abrir_hoja()
    log("Probando la conexion con el Google Sheets...")
    web.probar()

    log("Abriendo Chrome...")
    driver = crear_driver()
    primera = True
    try:
        esperar_sesion(driver, detener)
        log("Sesion lista.")
        for ini, fin in tramos(desde, hasta):
            if detener and detener.is_set():
                log("Detenido: lo ya cargado se queda en el Sheets.")
                break
            log(f"Bajando reporte del {ini:%d/%m/%Y} al {fin:%d/%m/%Y}...")
            encabezados, filas = leer_excel(descargar_tramo(driver, ini, fin))
            log(f"   {len(filas)} filas (rutas por dia).")
            if not encabezados:
                continue
            for j in range(0, max(len(filas), 1), FILAS_POR_ENVIO):
                parte = filas[j:j + FILAS_POR_ENVIO]
                r = web.cargar(encabezados, parte, ini.isoformat(),
                               fin.isoformat(), primera)
                primera = False
                resumen["filas"] += r.get("recibidas", 0)
                resumen["nuevas"] += r.get("nuevas", 0)
                resumen["actualizadas"] += r.get("actualizadas", 0)
                resumen["iguales"] += r.get("iguales", 0)
                resumen["total"] = r.get("total", resumen["total"])
                log(f"   Sheets: {r.get('nuevas', 0)} nuevas, "
                    f"{r.get('actualizadas', 0)} actualizadas, "
                    f"{r.get('iguales', 0)} sin cambio.")
    except Exception as e:
        resumen["error"] = str(e)
        log(f"ERROR: {e}")
    finally:
        soltar_driver(driver)
        log("Chrome queda abierto: cierralo tu cuando termines.")
    resumen["minutos"] = (time.time() - inicio) / 60
    return resumen


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        return
    desde = date.fromisoformat(sys.argv[1])
    hasta = date.fromisoformat(sys.argv[2])
    r = ejecutar(desde, hasta)
    print(json.dumps(r, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
